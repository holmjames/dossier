#!/usr/bin/env python3
"""Pull plain text out of a resume or job description file.

Usage:
    python3 extract_text.py <file> [--out text.txt]

Handles .docx, .pdf, .txt, .md, .html. Prints to stdout unless --out is given.
Tables in .docx files are flattened row by row so nothing hidden in a table
gets lost (ATS parsers often do lose it, which is its own warning).
"""
import argparse
import html
import re
import sys
from pathlib import Path


def from_docx(path: Path) -> str:
    try:
        import docx  # python-docx
    except ImportError:
        sys.exit("python-docx is not installed. Run: pip install python-docx")
    d = docx.Document(str(path))
    parts = []
    # Header text matters: many people put contact info there, and many ATS
    # parsers skip headers entirely.
    for section in d.sections:
        for p in section.header.paragraphs:
            if p.text.strip():
                parts.append(p.text)
    for p in d.paragraphs:
        parts.append(p.text)
    for t in d.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells]
            parts.append(" | ".join(c for c in cells if c))
    return "\n".join(parts)


def from_pdf(path: Path) -> str:
    try:
        import pdfplumber
        with pdfplumber.open(str(path)) as pdf:
            return "\n".join((page.extract_text() or "") for page in pdf.pages)
    except ImportError:
        pass
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except ImportError:
        sys.exit("Install pdfplumber or pypdf: pip install pdfplumber")


def from_html(path: Path) -> str:
    raw = path.read_text(errors="ignore")
    raw = re.sub(r"<(script|style)[^>]*>.*?</\1>", " ", raw, flags=re.S | re.I)
    raw = re.sub(r"<br\s*/?>|</p>|</li>|</h\d>|</div>", "\n", raw, flags=re.I)
    raw = re.sub(r"<[^>]+>", " ", raw)
    text = html.unescape(raw)
    return re.sub(r"[ \t]+", " ", text)


def extract(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".docx":
        text = from_docx(path)
    elif ext == ".pdf":
        text = from_pdf(path)
    elif ext in (".html", ".htm"):
        text = from_html(path)
    else:
        text = path.read_text(errors="ignore")
    # Normalise whitespace but keep line breaks, which the scorer uses for sections.
    text = re.sub(r"\r\n?", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip() + "\n"


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("file")
    ap.add_argument("--out", help="write text here instead of printing")
    args = ap.parse_args()
    path = Path(args.file)
    if not path.exists():
        sys.exit(f"No such file: {path}")
    text = extract(path)
    if args.out:
        Path(args.out).write_text(text)
        print(f"Wrote {len(text.split())} words to {args.out}")
    else:
        sys.stdout.write(text)


if __name__ == "__main__":
    main()
