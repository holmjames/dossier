#!/usr/bin/env python3
"""Keep the job-search tracker spreadsheet up to date.

Usage:
    python3 tracker.py add    --file tracker.xlsx --company "Adobe" --role "PM Intern" \
                              --location "San Jose, CA" --url "https://..." \
                              [--status "Drafted"] [--ats 78] [--notes "..."] \
                              [--resume "path"] [--cover "path"] [--source "Search"]
    python3 tracker.py update --file tracker.xlsx --company "Adobe" --role "PM Intern" \
                              --status "Applied" [--notes "..."] [--next "Follow up"] [--follow-up 2026-11-04]
    python3 tracker.py list   --file tracker.xlsx [--status Applied]

Creates the workbook if it does not exist. One row per application; the
company + role pair is the key. Statuses are a dropdown so the sheet stays
sortable: Sourced, Drafted, Approved, Applied, Screen, Interview, Offer,
Rejected, Withdrawn, Ghosted.
"""
import argparse
import sys
import re
from datetime import date, timedelta
from pathlib import Path

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:
    sys.exit("openpyxl is not installed. Run: pip install openpyxl")

COLUMNS = [
    ("Date added", 12), ("Company", 22), ("Role", 34), ("Location", 18),
    ("Source", 12), ("Status", 12), ("ATS score", 10), ("Applied on", 12),
    ("Next action", 26), ("Follow-up date", 14), ("Resume file", 36),
    ("Cover letter file", 36), ("Job URL", 50), ("Notes", 50),
]
STATUSES = ["Sourced", "Drafted", "Approved", "Applied", "Screen", "Interview",
            "Offer", "Rejected", "Withdrawn", "Ghosted"]
SHEET = "Applications"


def resolve_date(value):
    """Accept YYYY-MM-DD or +Nd (N days from today)."""
    if value is None:
        return None
    m = re.fullmatch(r"\+(\d+)d", value.strip())
    if m:
        return (date.today() + timedelta(days=int(m.group(1)))).isoformat()
    return value


def open_or_create(path: Path):
    if path.exists():
        wb = load_workbook(path)
        ws = wb[SHEET] if SHEET in wb.sheetnames else wb.active
        return wb, ws
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append([c for c, _ in COLUMNS])
    for i, (_, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F2937")
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}1"
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"F2:F1000")
    return wb, ws


def col(name: str) -> int:
    return [c for c, _ in COLUMNS].index(name) + 1


def find_row(ws, company: str, role: str):
    for r in range(2, ws.max_row + 1):
        c = (ws.cell(r, col("Company")).value or "").strip().lower()
        ro = (ws.cell(r, col("Role")).value or "").strip().lower()
        if c == company.strip().lower() and ro == role.strip().lower():
            return r
    return None


def cmd_add(args):
    path = Path(args.file)
    wb, ws = open_or_create(path)
    if find_row(ws, args.company, args.role):
        print(f"Already tracked: {args.company} / {args.role}. Use 'update' instead.")
        return
    row = {c: "" for c, _ in COLUMNS}
    row.update({
        "Date added": date.today().isoformat(),
        "Company": args.company, "Role": args.role, "Location": args.location or "",
        "Source": args.source or "", "Status": args.status or "Sourced",
        "ATS score": args.ats if args.ats is not None else "",
        "Applied on": date.today().isoformat() if (args.status or "") == "Applied" else "",
        "Next action": args.next or "", "Follow-up date": resolve_date(args.follow_up) or "",
        "Resume file": args.resume or "", "Cover letter file": args.cover or "",
        "Job URL": args.url or "", "Notes": args.notes or "",
    })
    ws.append([row[c] for c, _ in COLUMNS])
    r = ws.max_row
    if args.url:
        ws.cell(r, col("Job URL")).hyperlink = args.url
        ws.cell(r, col("Job URL")).style = "Hyperlink"
    wb.save(path)
    print(f"Added row {r}: {args.company} / {args.role} [{row['Status']}]")


def cmd_update(args):
    path = Path(args.file)
    if not path.exists():
        sys.exit(f"No tracker at {path}. Add a row first.")
    wb, ws = open_or_create(path)
    r = find_row(ws, args.company, args.role)
    if not r:
        sys.exit(f"Not found: {args.company} / {args.role}")
    changes = []
    for field, value in [("Status", args.status), ("ATS score", args.ats), ("Notes", args.notes),
                         ("Next action", args.next), ("Follow-up date", resolve_date(args.follow_up)),
                         ("Resume file", args.resume), ("Cover letter file", args.cover),
                         ("Location", args.location), ("Job URL", args.url)]:
        if value is not None:
            if field == "Notes" and ws.cell(r, col(field)).value and not args.replace_notes:
                value = f"{ws.cell(r, col(field)).value}\n{date.today().isoformat()}: {value}"
            ws.cell(r, col(field)).value = value
            changes.append(field)
    if args.status == "Applied" and not ws.cell(r, col("Applied on")).value:
        ws.cell(r, col("Applied on")).value = date.today().isoformat()
        changes.append("Applied on")
    wb.save(path)
    print(f"Updated row {r} ({args.company} / {args.role}): {', '.join(changes) or 'nothing'}")


def cmd_list(args):
    path = Path(args.file)
    if not path.exists():
        sys.exit(f"No tracker at {path}.")
    wb, ws = open_or_create(path)
    rows = []
    for r in range(2, ws.max_row + 1):
        status = ws.cell(r, col("Status")).value or ""
        if args.status and status.lower() != args.status.lower():
            continue
        rows.append((ws.cell(r, col("Date added")).value, ws.cell(r, col("Company")).value,
                     ws.cell(r, col("Role")).value, status, ws.cell(r, col("ATS score")).value,
                     ws.cell(r, col("Next action")).value))
    if not rows:
        print("No rows.")
        return
    print(f"{'Date':<11} {'Company':<20} {'Role':<34} {'Status':<10} {'ATS':<4} Next action")
    for d, c, ro, s, a, n in rows:
        print(f"{str(d or ''):<11} {str(c or '')[:19]:<20} {str(ro or '')[:33]:<34} {str(s or '')[:9]:<10} {str(a or ''):<4} {n or ''}")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    for name in ("add", "update", "list"):
        p = sub.add_parser(name)
        p.add_argument("--file", required=True)
        if name != "list":
            p.add_argument("--company", required=True)
            p.add_argument("--role", required=True)
            p.add_argument("--location")
            p.add_argument("--url")
            p.add_argument("--status", choices=STATUSES)
            p.add_argument("--ats", type=int)
            p.add_argument("--notes")
            p.add_argument("--next")
            p.add_argument("--follow-up", dest="follow_up", help="YYYY-MM-DD or +14d")
            p.add_argument("--resume")
            p.add_argument("--cover")
        if name == "add":
            p.add_argument("--source", help="Search, URL, Referral, Board, etc.")
        if name == "update":
            p.add_argument("--replace-notes", action="store_true", help="overwrite notes instead of appending")
        if name == "list":
            p.add_argument("--status")
    args = ap.parse_args()
    {"add": cmd_add, "update": cmd_update, "list": cmd_list}[args.cmd](args)


if __name__ == "__main__":
    main()
