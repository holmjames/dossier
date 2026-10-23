#!/usr/bin/env python3
"""Turn a plain-Markdown resume or cover letter into an ATS-safe Word file.

Usage:
    python3 build_docx.py resume.md  --kind resume --out "Jane Doe - Acme - PM - Resume.docx" [--pdf]
    python3 build_docx.py letter.md  --kind letter --out "Jane Doe - Acme - PM - Cover Letter.docx" [--pdf]

Options: --font "Calibri" --size 10.5 --margins 0.6 (inches) --name-size 16

Why a builder instead of writing Word files by hand: every output comes out as
one column, no tables, no text boxes, no header/footer, real bullets, and a
consistent look. Those are exactly the things ATS parsers need, and exactly
the things that drift when files are made ad hoc.

RESUME MARKDOWN CONVENTIONS (keep it this simple on purpose):

    # Jane Doe
    Austin, TX | jane@email.com | (512) 555-0100 | linkedin.com/in/janedoe

    Optional one- or two-line summary paragraph goes here.

    ## EXPERIENCE
    ### Acme Corp | Austin, TX
    *Senior Marketing Operations Manager* | Mar 2023 - Present
    - Bullet. Starts with a verb, ends with a number where possible.
    - Bullet.

    ### Previous Co | Remote
    *Marketing Operations Specialist* | Jun 2020 - Feb 2023
    - Bullet.

    ## EDUCATION
    ### University of Texas at Austin | Austin, TX
    *MBA, Class of 2028* | Expected May 2028
    - Optional bullet (honors, clubs, coursework).

    ## SKILLS & CERTIFICATIONS
    - Tools: Salesforce, Marketo, HubSpot, SQL, Tableau
    - Certifications: Salesforce Certified Administrator

A "### Org | Location" heading and an "*Title* | Dates" line (must start with *)
put the right side flush against the right margin (a tab stop, not a table).
Any other line containing "|" is rendered as plain text. "Skills: a, b, c" style
lines render as plain paragraphs with the label in bold.

COVER LETTER MARKDOWN: "# Name" and a contact line first, then plain paragraphs
separated by blank lines. A line starting with "Date:" is placed as the date.
"""
import argparse
import re
import shutil
import subprocess
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
except ImportError:
    sys.exit("python-docx is not installed. Run: pip install python-docx")


def set_base_style(doc, font, size):
    st = doc.styles["Normal"]
    st.font.name = font
    st.font.size = Pt(size)
    # East-Asian font mapping so Word does not substitute.
    rpr = st.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), font)
    pf = st.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0


def bottom_border(paragraph, size=6, color="000000"):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)
    ppr.append(pbdr)


def add_runs(paragraph, text, bold=False, italic=False, size=None):
    """Support **bold** and *italic* inline markers inside a line."""
    pattern = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")
    for chunk in pattern.split(text):
        if not chunk:
            continue
        b, i, t = bold, italic, chunk
        if chunk.startswith("**") and chunk.endswith("**"):
            b, t = True, chunk[2:-2]
        elif chunk.startswith("*") and chunk.endswith("*"):
            i, t = True, chunk[1:-1]
        run = paragraph.add_run(t)
        run.bold = b
        run.italic = i
        if size:
            run.font.size = Pt(size)


def two_sided(doc, left, right, right_margin_inches, bold_left=False, italic_left=False, italic_right=False):
    p = doc.add_paragraph()
    p.paragraph_format.tab_stops.add_tab_stop(Inches(right_margin_inches), WD_TAB_ALIGNMENT.RIGHT)
    add_runs(p, left.strip(), bold=bold_left, italic=italic_left)
    if right.strip():
        p.add_run("\t")
        add_runs(p, right.strip(), italic=italic_right)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.first_line_indent = Inches(-0.15)
    p.paragraph_format.space_after = Pt(1)
    add_runs(p, text)
    return p


def build_resume(lines, doc, text_width, name_size):
    i = 0
    # Name and contact block.
    while i < len(lines) and not lines[i].strip():
        i += 1
    if i < len(lines) and lines[i].startswith("# "):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(lines[i][2:].strip())
        r.bold = True
        r.font.size = Pt(name_size)
        i += 1
    while i < len(lines) and lines[i].strip() and not lines[i].startswith("#"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_runs(p, lines[i].strip())
        i += 1
    # Body.
    while i < len(lines):
        line = lines[i].rstrip()
        i += 1
        if not line.strip():
            continue
        if line.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(line[3:].strip().upper())
            r.bold = True
            bottom_border(p)
        elif line.startswith("### "):
            left, _, right = line[4:].partition("|")
            p = two_sided(doc, left, right, text_width, bold_left=True)
            p.paragraph_format.space_before = Pt(4)
        elif line.startswith("- ") or line.startswith("* "):
            bullet(doc, line[2:].strip())
        elif line.startswith("*") and "|" in line:
            left, _, right = line.partition("|")
            two_sided(doc, left, right, text_width, italic_right=False)
        elif re.match(r"^[A-Z][A-Za-z &/]{1,40}:\s", line):
            label, _, rest = line.partition(":")
            p = doc.add_paragraph()
            r = p.add_run(label + ":")
            r.bold = True
            add_runs(p, rest)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            add_runs(p, line.strip())


def build_letter(lines, doc, text_width, name_size):
    i = 0
    while i < len(lines) and not lines[i].strip():
        i += 1
    if i < len(lines) and lines[i].startswith("# "):
        p = doc.add_paragraph()
        r = p.add_run(lines[i][2:].strip())
        r.bold = True
        r.font.size = Pt(name_size)
        i += 1
        while i < len(lines) and lines[i].strip() and not lines[i].startswith("#"):
            p = doc.add_paragraph()
            add_runs(p, lines[i].strip())
            i += 1
        doc.add_paragraph()
    # Each line is its own paragraph; a blank line adds space after the block,
    # so "Sincerely," and the name stay on separate lines.
    block = []

    def flush():
        for n, text in enumerate(block):
            text = text.strip()
            if text.lower().startswith("date:"):
                text = text[5:].strip()
            if text.startswith("- "):
                bullet(doc, text[2:])
                continue
            p = doc.add_paragraph()
            if n == len(block) - 1:
                p.paragraph_format.space_after = Pt(10)
            add_runs(p, text)
        block.clear()

    for line in lines[i:]:
        if not line.strip():
            flush()
        else:
            block.append(line)
    flush()


def to_pdf(docx_path: Path):
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        print("PDF skipped: LibreOffice not found (open the .docx in Word and export, or install LibreOffice).")
        return None
    subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", str(docx_path.parent), str(docx_path)],
                   check=False, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
    pdf = docx_path.with_suffix(".pdf")
    if pdf.exists():
        print(f"Wrote {pdf}")
        return pdf
    print("PDF conversion did not produce a file.")
    return None


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("markdown")
    ap.add_argument("--kind", choices=["resume", "letter"], required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--font", default="Calibri")
    ap.add_argument("--size", type=float, help="point size (default 10.5 resume, 11 letter)")
    ap.add_argument("--name-size", type=float, default=16)
    ap.add_argument("--margins", type=float, help="inches, all sides (default 0.6 resume, 1.0 letter)")
    ap.add_argument("--pdf", action="store_true", help="also export a PDF via LibreOffice if available")
    args = ap.parse_args()

    if args.size is None:
        args.size = 10.5 if args.kind == "resume" else 11
    if args.margins is None:
        args.margins = 0.6 if args.kind == "resume" else 1.0
    lines = Path(args.markdown).read_text().split("\n")
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)  # US Letter
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(args.margins))
    text_width = 8.5 - 2 * args.margins
    set_base_style(doc, args.font, args.size)

    if args.kind == "resume":
        build_resume(lines, doc, text_width, args.name_size)
    else:
        build_letter(lines, doc, text_width, args.name_size)

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    print(f"Wrote {out}")
    if args.pdf:
        to_pdf(out)


if __name__ == "__main__":
    main()
